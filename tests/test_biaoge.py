# -*- coding: utf-8 -*-
"""biaoge CLI 冒烟测试（pytest）"""
import subprocess
import sys
import os
from pathlib import Path

HERE = Path(__file__).parent
SRC = HERE.parent / "src" / "biaoge.py"
PY = sys.executable


def run(*args, cwd=None):
    r = subprocess.run([PY, str(SRC), *args], capture_output=True, text=True, cwd=cwd)
    assert r.returncode == 0, f"命令失败: {args}\nSTDOUT: {r.stdout}\nSTDERR: {r.stderr}"
    return r.stdout


def make_sample(dirpath: Path):
    """造两个 csv 样本"""
    dirpath.mkdir(parents=True, exist_ok=True)
    (dirpath / "1月.csv").write_text("姓名,部门,工资\n张三,销售,8000\n李四,技术,12000\n", encoding="utf-8")
    (dirpath / "2月.csv").write_text("姓名,部门,工资\n张三,销售,8500\n王五,行政,6000\n", encoding="utf-8")
    return dirpath


def test_summary(tmp_path):
    d = make_sample(tmp_path / "报表")
    out = run("summary", str(d))
    assert "1月.csv" in out and "2月.csv" in out


def test_merge(tmp_path):
    d = make_sample(tmp_path / "报表")
    out_file = tmp_path / "合并.xlsx"
    run("merge", str(d), "-o", str(out_file), "--source-col")
    assert out_file.exists() and out_file.stat().st_size > 0
    import openpyxl
    ws = openpyxl.load_workbook(out_file).active
    rows = list(ws.values)
    assert rows[0][0] == "姓名"
    assert len(rows) == 5  # 表头 + 4 行


def test_dedup(tmp_path):
    src = tmp_path / "dup.csv"
    src.write_text("手机号,姓名\n13800000000,A\n13800000000,A\n13900000000,B\n", encoding="utf-8")
    out_file = tmp_path / "clean.csv"
    run("dedup", str(src), "--columns", "手机号", "-o", str(out_file))
    lines = out_file.read_text(encoding="utf-8").strip().splitlines()
    assert len(lines) == 3  # 表头 + 2 行去重后


def test_convert(tmp_path):
    src = tmp_path / "a.csv"
    src.write_text("x,y\n1,2\n", encoding="utf-8")
    outdir = tmp_path / "out"
    run("convert", str(src), "-d", str(outdir), "--fmt", ".xlsx")
    assert any(p.suffix == ".xlsx" for p in outdir.iterdir())


def test_split(tmp_path):
    src = tmp_path / "staff.csv"
    src.write_text("姓名,部门\n张三,销售\n李四,技术\n王五,销售\n", encoding="utf-8")
    outdir = tmp_path / "split"
    run("split", str(src), "--column", "部门", "-d", str(outdir))
    files = list(outdir.rglob("*"))
    assert len([f for f in files if f.is_file()]) >= 2


def test_join(tmp_path):
    left = tmp_path / "订单.csv"
    left.write_text("订单号,客户编号,金额\nD1,C001,100\nD2,C002,250\nD3,C999,80\n", encoding="utf-8")
    right = tmp_path / "客户.csv"
    right.write_text("客户编号,姓名,城市\nC001,张三,北京\nC002,李四,上海\n", encoding="utf-8")
    out = tmp_path / "join.xlsx"
    run("join", str(left), str(right), "--on", "客户编号", "-o", str(out))
    import openpyxl
    ws = openpyxl.load_workbook(out).active
    rows = list(ws.values)
    assert rows[0] == ("订单号", "客户编号", "金额", "姓名", "城市")
    assert rows[1][3] == "张三"
    assert rows[3][3] in ("", None)  # 未匹配留空


def test_join_only_matched(tmp_path):
    left = tmp_path / "a.csv"
    left.write_text("k,v\n1,x\n2,y\n9,z\n", encoding="utf-8")
    right = tmp_path / "b.csv"
    right.write_text("k,name\n1,A\n2,B\n", encoding="utf-8")
    out = tmp_path / "j.csv"
    run("join", str(left), str(right), "--on", "k", "--only-matched", "-o", str(out))
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert len(lines) == 3  # 表头+2


def test_groupby_default_sum(tmp_path):
    src = tmp_path / "sales.csv"
    src.write_text(
        "产品,销售额,数量\nA,100,2\nB,200,3\nA,50,1\nB,20,4\n",
        encoding="utf-8")
    out = tmp_path / "g.csv"
    run("groupby", str(src), "--by", "产品", "-o", str(out))
    text = out.read_text(encoding="utf-8")
    lines = text.strip().splitlines()
    assert len(lines) == 3  # 表头 + A + B
    assert "sum(销售额)" in lines[0]
    # A 组销售额合计 150
    assert "150" in lines[1]


def test_groupby_multi_agg_and_sort(tmp_path):
    src = tmp_path / "orders.csv"
    src.write_text(
        "客户,金额\n张三,100\n李四,30\n张三,50\n王五,80\n",
        encoding="utf-8")
    out = tmp_path / "agg.xlsx"
    run("groupby", str(src), "--by", "客户", "--agg", "金额:sum,金额:count,金额:avg",
        "--sort", "金额", "-o", str(out))
    import openpyxl
    ws = openpyxl.load_workbook(out).active
    rows = list(ws.values)
    assert rows[0] == ("客户", "sum(金额)", "count(金额)", "avg(金额)")
    # 默认降序：张三(150) 第一
    assert rows[1][0] == "张三"
    assert rows[1][1] == 150
    assert rows[1][2] == 2
    assert rows[1][3] == 75.0


def test_filter_numeric_and_contains(tmp_path):
    src = tmp_path / "orders.csv"
    src.write_text(
        "地区,销售额,状态\n华东,1200,已完成\n华南,800,已完成\n华东,300,退款\n华北,5000,已完成\n",
        encoding="utf-8")
    out = tmp_path / "f.csv"
    run("filter", str(src), "--where", "销售额>500", "--where", "地区~华东", "-o", str(out))
    text = out.read_text(encoding="utf-8").strip()
    lines = text.splitlines()
    assert len(lines) == 2  # 表头 + 华东1200
    assert "1200" in lines[1]


def test_filter_neq_and_dropempty(tmp_path):
    src = tmp_path / "a.csv"
    src.write_text(
        "k,v\n1,x\n2,退款\n3,\n,,\n",
        encoding="utf-8")
    out = tmp_path / "f2.csv"
    run("filter", str(src), "--where", "v!=退款", "--drop-empty", "-o", str(out))
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    # 保留：1,x 和 3,（k=3 非全空行）；干掉：退款行和 ,, 全空行
    assert len(lines) == 3
    assert lines[1].startswith("1,x")
    assert lines[2].startswith("3,")


def test_filter_bad_column(tmp_path):
    src = tmp_path / "a.csv"
    src.write_text("k,v\n1,x\n", encoding="utf-8")
    out = tmp_path / "f3.csv"
    r = subprocess.run([PY, str(SRC), "filter", str(src), "--where", "不存在=1", "-o", str(out)],
                       capture_output=True, text=True)
    assert r.returncode == 2
    assert "没有列" in r.stdout


def test_clean_spaces_dates_fill(tmp_path):
    src = tmp_path / "dirty.csv"
    # 全角空格 + 首尾空格 + 混乱日期 + 缺失值
    src.write_text(
        "姓名,入职日期,备注\n\u3000张三 ,2026/3/5,正式\n李四,2026年1月15日,\n\n,,\n",
        encoding="utf-8")
    out = tmp_path / "clean.xlsx"
    run("clean", str(src), "--dates", "--fill", "-", "-o", str(out))
    import openpyxl
    ws = openpyxl.load_workbook(out).active
    rows = [r for r in ws.values]
    assert rows[0] == ("姓名", "入职日期", "备注")
    assert rows[1][0] == "张三"          # 去空格
    assert rows[1][1] == "2026-03-05"    # 日期统一
    assert rows[2][1] == "2026-01-15"    # 中文日期
    assert rows[2][2] == "-"             # 填缺失
    assert len(rows) == 3                # 空行被去掉


def test_clean_only_columns(tmp_path):
    src = tmp_path / "d2.csv"
    src.write_text("a,b\n  x  ,y  y\n", encoding="utf-8")
    out = tmp_path / "c2.csv"
    run("clean", str(src), "--columns", "a", "-o", str(out))
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    cells = lines[1].split(",")
    assert cells[0] == "x"        # a 列被清洗
    assert cells[1] == "y  y"     # b 列保持原样（内部空格不动）


# ---------- report 数据体检 ----------
def test_report_basic(tmp_path):
    src = tmp_path / "客户表.csv"
    src.write_text(
        "姓名,手机号,消费额,注册日期\n"
        "张三,13800138000,1200,2026-01-05\n"
        "李四,13900139000,,2026/2/10\n"
        "张三,13800138000,1200,2026-01-05\n"   # 完全重复行
        "\n"                                     # 空行
        " 王五 ,13700137000,99999,2026年3月1日\n",  # 首尾空格+离群值+中文日期
        encoding="utf-8")
    out = run("report", str(src))
    # 结构断言
    assert "# 数据体检报告" in out
    assert "行数：**5**" in out and "列数：**4**" in out
    # 检测到完全重复行
    assert "完全重复的行：**1**" in out
    # 检测到全空行
    assert "全空行：**1**" in out
    # 手机号列类型识别 + 隐私提示
    assert "疑似手机号" in out
    # 消费额离群值提示（99999 远大于均值）
    assert "离群值" in out
    # 注册日期格式不统一
    assert "日期格式不统一" in out
    # 首尾空格
    assert "首尾带空格" in out
    # 总评
    assert "总评：" in out


def test_report_to_file(tmp_path):
    src = tmp_path / "d.csv"
    src.write_text("a,b\n1,x\n2,y\n", encoding="utf-8")
    rep = tmp_path / "out" / "体检.md"
    out = run("report", str(src), "-o", str(rep))
    assert rep.exists()
    content = rep.read_text(encoding="utf-8")
    assert "# 数据体检报告" in content
    assert "行数：**2**" in content
    # 干净数据高分
    assert "分 ——" in content


def test_report_primary_key(tmp_path):
    src = tmp_path / "k.csv"
    src.write_text("编号,名称\n1,a\n2,b\n3,c\n", encoding="utf-8")
    out = run("report", str(src))
    # 编号列全唯一 → 主键提示
    assert "可作主键" in out


def test_report_missing_header(tmp_path):
    src = tmp_path / "empty.csv"
    src.write_text("", encoding="utf-8")
    import subprocess
    r = subprocess.run([PY, str(SRC), "report", str(src)], capture_output=True, text=True)
    assert r.returncode == 2
    assert "读不到表头" in r.stdout


# ---------- watch（v0.6） ----------
def _watch_once(folder, out, *extra):
    return run("watch", str(folder), "--out", str(out), "--once", *extra)


def test_watch_once_basic(tmp_path):
    d = make_sample(tmp_path / "收集箱")
    out = tmp_path / "总表.xlsx"
    outp = _watch_once(d, out)
    assert "已汇总 2 个文件" in outp
    import openpyxl
    ws = openpyxl.load_workbook(out).active
    rows = list(ws.values)
    assert rows[0][-1] == "来源文件"
    assert len(rows) == 5  # 表头 + 4 行


def test_watch_once_no_change(tmp_path):
    d = make_sample(tmp_path / "收集箱")
    out = tmp_path / "总表.xlsx"
    _watch_once(d, out)
    outp2 = _watch_once(d, out)
    assert "无变化" in outp2


def test_watch_picks_up_new_file(tmp_path):
    d = make_sample(tmp_path / "收集箱")
    out = tmp_path / "总表.xlsx"
    _watch_once(d, out)
    # 新丢一个文件进收集箱
    (d / "3月.csv").write_text("姓名,部门,工资\n赵六,财务,9000\n", encoding="utf-8")
    outp = _watch_once(d, out)
    assert "已汇总 3 个文件" in outp
    import openpyxl
    ws = openpyxl.load_workbook(out).active
    rows = list(ws.values)
    assert len(rows) == 6  # 表头 + 5 行
    assert any(r[-1] == "3月.csv" for r in rows[1:])


def test_watch_skips_output_temp_hidden(tmp_path):
    d = make_sample(tmp_path / "收集箱")
    out = d / "总表.xlsx"  # 输出放被监控目录内，应自动跳过
    (d / "~$临时文件.xlsx").write_bytes(b"junk")   # Office 临时文件
    (d / ".隐藏.csv").write_text("a,b\n1,2\n", encoding="utf-8")  # 隐藏文件
    outp = _watch_once(d, out)
    assert "已汇总 2 个文件" in outp  # 只有 1月/2月，其余全跳过


def test_watch_live_monitor(tmp_path):
    """真实守护模式：后台起 watch，丢新文件，验证自动重建"""
    import subprocess, time
    d = make_sample(tmp_path / "收集箱")
    out = tmp_path / "总表.xlsx"
    proc = subprocess.Popen(
        [PY, str(SRC), "watch", str(d), "--out", str(out), "--interval", "1"],
        stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
    try:
        time.sleep(2.5)  # 等首轮汇总
        (d / "4月.csv").write_text("姓名,部门,工资\n孙七,运营,7000\n", encoding="utf-8")
        time.sleep(3.0)  # 等检测
        import openpyxl
        ws = openpyxl.load_workbook(out).active
        rows = list(ws.values)
        assert len(rows) == 6, f"实时模式未拾取新文件，rows={len(rows)}"
        assert any(r[-1] == "4月.csv" for r in rows[1:])
    finally:
        proc.terminate(); proc.wait(timeout=5)
        log = proc.stdout.read()
        assert "监控中" in log


# ---------- pivot（v0.7） ----------
def _pivot_sample(tmp_path):
    p = tmp_path / "订单.csv"
    p.write_text(
        "地区,月份,产品,金额\n"
        "华东,1月,A,100\n"
        "华东,1月,B,50\n"
        "华东,2月,A,200\n"
        "华北,1月,A,80\n"
        "华北,2月,B,120\n"
        "华南,1月,B,60\n"
        "华南,2月,A,90\n"
        "华北,1月,C,70\n",
        encoding="utf-8")
    return p


def test_pivot_sum_margins(tmp_path):
    src = _pivot_sample(tmp_path)
    out = tmp_path / "pv.xlsx"
    run("pivot", str(src), "--rows", "地区", "--cols", "月份",
        "--values", "金额:sum", "--margins", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    header = rows[0]
    assert header[0] == "地区" and "1月" in header and "2月" in header and header[-1] == "总计"
    d = {r[0]: r for r in rows[1:]}
    assert d["华东"][1] == 150 and d["华东"][2] == 200 and d["华东"][-1] == 350
    assert d["总计"][1] == 360 and d["总计"][-1] == 770
    # 默认按合计降序：华东(350) > 华北(270) > 华南(150) > 总计行在最后
    assert rows[1][0] == "华东" and rows[2][0] == "华北" and rows[3][0] == "华南"


def test_pivot_count(tmp_path):
    src = _pivot_sample(tmp_path)
    out = tmp_path / "pc.xlsx"
    run("pivot", str(src), "--rows", "地区", "--cols", "产品",
        "--values", "金额:count", "--margins", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    d = {r[0]: r for r in rows[1:]}
    header = rows[0]
    ia, ib, ic, it = header.index("A"), header.index("B"), header.index("C"), header.index("总计")
    assert d["华东"][ia] == 2 and d["华东"][ib] == 1 and d["华东"][ic] in (None, "")
    assert d["总计"][ia] == 4 and d["总计"][ib] == 3 and d["总计"][ic] == 1 and d["总计"][it] == 8


def test_pivot_multi_rows_keep_order(tmp_path):
    src = _pivot_sample(tmp_path)
    out = tmp_path / "pm.xlsx"
    run("pivot", str(src), "--rows", "地区,产品", "--cols", "月份",
        "--values", "金额", "--keep-order", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    assert rows[0][0] == "地区" and rows[0][1] == "产品"
    # keep-order：行按原始出现顺序
    assert rows[1][0] == "华东" and rows[1][1] == "A"
    assert rows[2][0] == "华东" and rows[2][1] == "B"
    assert len(rows) == 8  # 表头 + 7 组合


def test_pivot_avg(tmp_path):
    src = _pivot_sample(tmp_path)
    out = tmp_path / "pa.xlsx"
    run("pivot", str(src), "--rows", "地区", "--cols", "产品",
        "--values", "金额:avg", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    header = rows[0]
    ia = header.index("A")
    d = {r[0]: r for r in rows[1:]}
    assert d["华东"][ia] == 150  # (100+200)/2


def test_pivot_bad_column(tmp_path):
    src = _pivot_sample(tmp_path)
    r = subprocess.run([PY, str(SRC), "pivot", str(src), "--rows", "不存在列",
                        "--cols", "月份", "--values", "金额", "-o", str(tmp_path / "x.xlsx")],
                       capture_output=True, text=True)
    assert r.returncode == 2
    assert "行字段" in r.stdout


def test_pivot_bad_func(tmp_path):
    src = _pivot_sample(tmp_path)
    r = subprocess.run([PY, str(SRC), "pivot", str(src), "--rows", "地区",
                        "--cols", "月份", "--values", "金额:median", "-o", str(tmp_path / "x.xlsx")],
                       capture_output=True, text=True)
    assert r.returncode == 2
    assert "聚合函数" in r.stdout


def test_pivot_nonnumeric_values(tmp_path):
    p = tmp_path / "t.csv"
    p.write_text("a,b,v\nx,y,文本1\nx,z,文本2\n", encoding="utf-8")
    out = tmp_path / "t.xlsx"
    run("pivot", str(p), "--rows", "a", "--cols", "b", "--values", "v:sum", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    assert rows[1][0] == "x"
    # 非数值进不了 sum 桶 → 单元格留空（与 Excel 透视表行为一致），关键是没崩溃
    assert rows[1][1] in (None, "", 0) and rows[1][2] in (None, "", 0)


def test_pivot_blank_cells(tmp_path):
    """列字段有空值 → 归入 (空白) 列"""
    p = tmp_path / "b.csv"
    p.write_text("地区,月份,金额\n华东,,100\n华东,1月,200\n", encoding="utf-8")
    out = tmp_path / "b.xlsx"
    run("pivot", str(p), "--rows", "地区", "--cols", "月份", "--values", "金额", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    assert "(空白)" in rows[0]


# ---------- unpivot 宽表转长表 ----------
def _unpivot_sample(tmp_path):
    p = tmp_path / "宽表.csv"
    p.write_text("门店,区域,1月,2月,3月\n"
                 "一店,华东,100,200,300\n"
                 "二店,华南,150,,250\n", encoding="utf-8")
    return p


def test_unpivot_basic(tmp_path):
    src = _unpivot_sample(tmp_path)
    out = tmp_path / "长表.xlsx"
    run("unpivot", str(src), "--id", "门店,区域", "--name", "月份", "--value", "销售额", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    assert list(rows[0]) == ["门店", "区域", "月份", "销售额"]
    # 2行 × 3个月 - 1个空值 = 5 数据行（+表头=6）
    assert len(rows) == 6
    assert rows[1][2] == "1月" and float(rows[1][3]) == 100
    assert rows[4][0] == "二店" and rows[4][2] == "1月"
    # 二店2月为空被跳过 → 直接到3月
    assert rows[5][0] == "二店" and rows[5][2] == "3月"


def test_unpivot_skip_and_keep_empty(tmp_path):
    src = _unpivot_sample(tmp_path)
    out1 = tmp_path / "skip.xlsx"
    r = run("unpivot", str(src), "--id", "门店", "-o", str(out1))
    assert "跳过空值" in r
    out2 = tmp_path / "keep.xlsx"
    run("unpivot", str(src), "--id", "门店", "--keep-empty", "-o", str(out2))
    import openpyxl
    rows = list(openpyxl.load_workbook(out2).active.values)
    # --id 只有门店 → 拆4列(区域,1月,2月,3月)，2行×4=8数据行+表头=9
    assert len(rows) == 9
    # skip 版少一行（空2月被跳过）
    import openpyxl as _o
    rows1 = list(_o.load_workbook(out1).active.values)
    assert len(rows1) == 8


def test_unpivot_auto_cols(tmp_path):
    """不给 --cols：除 ID 外全部列自动拆"""
    src = _unpivot_sample(tmp_path)
    out = tmp_path / "auto.xlsx"
    run("unpivot", str(src), "--id", "门店,区域", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    names = {r[2] for r in rows[1:]}
    assert names == {"1月", "2月", "3月"}


def test_unpivot_no_id(tmp_path):
    """不给 --id：全部列都拆"""
    src = tmp_path / "n.csv"
    src.write_text("甲,乙\n1,2\n", encoding="utf-8")
    out = tmp_path / "n.xlsx"
    run("unpivot", str(src), "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    assert list(rows[0]) == ["指标", "值"]
    assert len(rows) == 3  # 1数据行×2列+表头


def test_unpivot_bad_col(tmp_path):
    src = _unpivot_sample(tmp_path)
    r = subprocess.run([PY, str(SRC), "unpivot", str(src), "--id", "不存在",
                        "-o", str(tmp_path / "x.xlsx")], capture_output=True, text=True)
    assert r.returncode == 2
    assert "ID列" in r.stdout


def test_unpivot_pivot_roundtrip(tmp_path):
    """unpivot → pivot 回环：拆完再聚合回去，数值应还原"""
    src = tmp_path / "rt.csv"
    src.write_text("门店,1月,2月\n"
                   "A,10,20\nB,30,40\n", encoding="utf-8")
    mid = tmp_path / "mid.csv"
    run("unpivot", str(src), "--id", "门店", "--name", "月份", "--value", "金额", "-o", str(mid))
    out = tmp_path / "rt2.xlsx"
    run("pivot", str(mid), "--rows", "门店", "--cols", "月份", "--values", "金额:sum",
        "--keep-order", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    hdr = list(rows[0])
    body = {r[0]: r for r in rows[1:]}
    i1, i2 = hdr.index("1月"), hdr.index("2月")
    assert float(body["A"][i1]) == 10 and float(body["A"][i2]) == 20
    assert float(body["B"][i1]) == 30 and float(body["B"][i2]) == 40


# ---------- diff 两表对比（v0.9） ----------
def _diff_sample(tmp_path):
    old = tmp_path / "old.csv"
    old.write_text("编号,姓名,工资\nA001,张三,8000\nA002,李四,12000\nA003,王五,9000\n", encoding="utf-8")
    new = tmp_path / "new.csv"
    new.write_text("编号,姓名,工资\nA001,张三,8500\nA003,王五,9000\nA004,赵六,7000\n", encoding="utf-8")
    return old, new


def test_diff_basic_three_categories(tmp_path):
    """基础三类差异：A001 修改，A002 删除，A004 新增"""
    old, new = _diff_sample(tmp_path)
    out = run("diff", str(old), str(new), "--key", "编号")
    assert "新增：1" in out
    assert "删除：1" in out
    assert "修改：1" in out
    assert "A004" in out  # 新增
    assert "A002" in out  # 删除
    assert "8000" in out and "8500" in out  # 修改明细


def test_diff_custom_columns(tmp_path):
    """--columns 只对比指定列"""
    old = tmp_path / "o.csv"
    old.write_text("k,a,b\n1,x,old\n2,y,ok\n", encoding="utf-8")
    new = tmp_path / "n.csv"
    new.write_text("k,a,b\n1,x,new\n2,changed,ok\n", encoding="utf-8")
    out = run("diff", str(old), str(new), "--key", "k", "--columns", "b")
    # 只对比 b 列：k=1 的 b 变了，k=2 的 a 变了但不看
    assert "修改：1" in out


def test_diff_key_not_found(tmp_path):
    """key 列不存在 → 友好报错 exit=1"""
    old, new = _diff_sample(tmp_path)
    r = subprocess.run([PY, str(SRC), "diff", str(old), str(new), "--key", "不存在"],
                       capture_output=True, text=True)
    assert r.returncode == 1
    assert "没有关键列" in r.stdout


def test_diff_number_normalization(tmp_path):
    """数字归一：1.0 == 1 == '1' 不误报"""
    old = tmp_path / "o.csv"
    old.write_text("k,v\nA,1\nB,2.5\n", encoding="utf-8")
    new = tmp_path / "n.csv"
    new.write_text("k,v\nA,1.0\nB,2.5\n", encoding="utf-8")
    out = run("diff", str(old), str(new), "--key", "k")
    # 1 == 1.0 不应报修改
    assert "修改：0" in out


def test_diff_key_case_and_space(tmp_path):
    """key 匹配：忽略大小写和首尾空格"""
    old = tmp_path / "o.csv"
    old.write_text("k,v\n  ABC ,1\nDEF,2\n", encoding="utf-8")
    new = tmp_path / "n.csv"
    new.write_text("k,v\nabc,10\ndef,20\n", encoding="utf-8")
    out = run("diff", str(old), str(new), "--key", "k")
    # ABC 和 abc 应匹配，DEF 和 def 应匹配 → 都是修改
    assert "新增：0" in out
    assert "删除：0" in out
    assert "修改：2" in out


def test_diff_output_xlsx_sheets(tmp_path):
    """-o 输出 xlsx，验证 4 个 sheet 结构"""
    old, new = _diff_sample(tmp_path)
    out_xlsx = tmp_path / "diff.xlsx"
    run("diff", str(old), str(new), "--key", "编号", "-o", str(out_xlsx))
    import openpyxl
    wb = openpyxl.load_workbook(out_xlsx)
    assert wb.sheetnames == ["摘要", "新增", "删除", "修改"]
    # 摘要 sheet
    ws_sum = wb["摘要"]
    assert ws_sum["A1"].value == "项目"
    assert ws_sum["B2"].value == 3  # 旧表行数
    assert ws_sum["B3"].value == 3  # 新表行数
    assert ws_sum["B4"].value == 1  # 新增
    assert ws_sum["B5"].value == 1  # 删除
    # 修改 sheet 结构：key, 列名, 旧值, 新值
    ws_mod = wb["修改"]
    mod_header = [c.value for c in ws_mod[1]]
    assert mod_header == ["编号", "列名", "旧值", "新值"]
    # A001 工资变了
    mod_rows = list(ws_mod.iter_rows(min_row=2, values_only=True))
    assert len(mod_rows) == 1
    assert mod_rows[0][0] == "A001"
    assert mod_rows[0][1] == "工资"


def test_diff_no_common_columns(tmp_path):
    """两表列完全不同 → 友好提示"""
    old = tmp_path / "o.csv"
    old.write_text("k,a\n1,x\n", encoding="utf-8")
    new = tmp_path / "n.csv"
    new.write_text("k,b\n1,y\n", encoding="utf-8")
    r = subprocess.run([PY, str(SRC), "diff", str(old), str(new), "--key", "k"],
                       capture_output=True, text=True)
    assert r.returncode == 1
    assert "没有共同列" in r.stdout


# ---------- merge join（v0.10） ----------
def _merge_join_sample(tmp_path):
    left = tmp_path / "left.csv"
    left.write_text("工号,姓名,部门\n001,张三,销售\n002,李四,技术\n003,王五,市场\n", encoding="utf-8")
    right = tmp_path / "right.csv"
    right.write_text("工号,工资,入职年份\n001,8000,2020\n002,12000,2019\n004,9000,2021\n", encoding="utf-8")
    return left, right


def test_merge_join_left(tmp_path):
    """left join：左表全保留，右表匹配不上的填空"""
    left, right = _merge_join_sample(tmp_path)
    out = tmp_path / "out.xlsx"
    out_text = run("merge", str(left), str(right), "--on", "工号", "--how", "left", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    assert rows[0] == ("工号", "姓名", "部门", "工资", "入职年份")
    assert len(rows) == 4  # 表头 + 3 行（左表全保留）
    d = {r[0]: r for r in rows[1:]}
    assert d["001"][3] == 8000 or d["001"][3] == "8000"
    assert d["003"][3] in (None, "")  # 王五在右表无匹配
    assert "匹配" in out_text and "左表" in out_text


def test_merge_join_inner(tmp_path):
    """inner join：只保留两表都匹配的 key"""
    left, right = _merge_join_sample(tmp_path)
    out = tmp_path / "out.xlsx"
    run("merge", str(left), str(right), "--on", "工号", "--how", "inner", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    assert len(rows) == 3  # 表头 + 2 行（001, 002）
    keys = {r[0] for r in rows[1:]}
    assert keys == {"001", "002"}


def test_merge_join_right(tmp_path):
    """right join：右表全保留，左表匹配不上的填空"""
    left, right = _merge_join_sample(tmp_path)
    out = tmp_path / "out.xlsx"
    run("merge", str(left), str(right), "--on", "工号", "--how", "right", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    assert len(rows) == 4  # 表头 + 3 行（右表全保留）
    d = {r[0]: r for r in rows[1:]}
    assert d["004"][1] in (None, "")  # 004 在左表无匹配，姓名为空
    assert d["004"][2] in (None, "")  # 部门也为空
    assert d["001"][1] == "张三"


def test_merge_join_multi_key(tmp_path):
    """多列 key 关联"""
    left = tmp_path / "l.csv"
    left.write_text("部门,工号,姓名\n销售,001,张三\n销售,002,李四\n技术,001,王五\n", encoding="utf-8")
    right = tmp_path / "r.csv"
    right.write_text("部门,工号,工资\n销售,001,8000\n技术,001,10000\n", encoding="utf-8")
    out = tmp_path / "out.csv"
    run("merge", str(left), str(right), "--on", "部门,工号", "-o", str(out))
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert len(lines) == 4  # 表头 + 3 行
    assert "工资" in lines[0]
    # 销售,001 → 8000；技术,001 → 10000；销售,002 → 空
    assert "8000" in lines[1]
    assert "10000" in lines[3] or "10000" in lines[2]


def test_merge_join_number_normalization(tmp_path):
    """数字 key 归一：1.0 == 1 == '1'"""
    left = tmp_path / "l.csv"
    left.write_text("k,v\n1,a\n2,b\n", encoding="utf-8")
    right = tmp_path / "r.csv"
    right.write_text("k,w\n1.0,x\n2,y\n", encoding="utf-8")
    out = tmp_path / "out.csv"
    out_text = run("merge", str(left), str(right), "--on", "k", "-o", str(out))
    assert "匹配率：100.0%" in out_text
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert len(lines) == 3  # 表头 + 2 行全匹配


def test_merge_join_key_case_and_space(tmp_path):
    """key 归一：忽略大小写和首尾空格"""
    left = tmp_path / "l.csv"
    left.write_text("k,v\n  ABC ,1\nDEF,2\n", encoding="utf-8")
    right = tmp_path / "r.csv"
    right.write_text("k,w\nabc,10\ndef,20\n", encoding="utf-8")
    out = tmp_path / "out.csv"
    out_text = run("merge", str(left), str(right), "--on", "k", "-o", str(out))
    assert "匹配率：100.0%" in out_text


def test_merge_join_suffix_conflict(tmp_path):
    """同名列冲突：右表重名列加后缀"""
    left = tmp_path / "l.csv"
    left.write_text("k,a,b\n1,x,L1\n", encoding="utf-8")
    right = tmp_path / "r.csv"
    right.write_text("k,a,c\n1,R1,R2\n", encoding="utf-8")
    out = tmp_path / "out.xlsx"
    run("merge", str(left), str(right), "--on", "k", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    # 列名：k, a, b, a_右, c
    assert rows[0] == ("k", "a", "b", "a_右", "c")
    assert rows[1][1] == "x"      # 左表 a
    assert rows[1][3] == "R1"     # 右表 a → a_右


def test_merge_join_dup_first(tmp_path):
    """dup=first：右表 key 重复取第一条"""
    left = tmp_path / "l.csv"
    left.write_text("k,v\n1,a\n", encoding="utf-8")
    right = tmp_path / "r.csv"
    right.write_text("k,w\n1,first\n1,second\n", encoding="utf-8")
    out = tmp_path / "out.csv"
    out_text = run("merge", str(left), str(right), "--on", "k", "--dup", "first", "-o", str(out))
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert "first" in lines[1]
    assert "second" not in lines[1]
    assert "重复" in out_text


def test_merge_join_dup_error(tmp_path):
    """dup=error：右表 key 重复报错 exit 1"""
    left = tmp_path / "l.csv"
    left.write_text("k,v\n1,a\n", encoding="utf-8")
    right = tmp_path / "r.csv"
    right.write_text("k,w\n1,first\n1,second\n", encoding="utf-8")
    out = tmp_path / "out.csv"
    r = subprocess.run([PY, str(SRC), "merge", str(left), str(right), "--on", "k", "--dup", "error", "-o", str(out)],
                       capture_output=True, text=True)
    assert r.returncode == 1
    assert "重复" in r.stdout


def test_merge_join_dup_concat(tmp_path):
    """dup=concat：文本列拼接，数字列取首条"""
    left = tmp_path / "l.csv"
    left.write_text("k,v\n1,a\n", encoding="utf-8")
    right = tmp_path / "r.csv"
    right.write_text("k,w,num\n1,hello,100\n1,world,200\n", encoding="utf-8")
    out = tmp_path / "out.csv"
    run("merge", str(left), str(right), "--on", "k", "--dup", "concat", "-o", str(out))
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert "hello|world" in lines[1]  # 文本拼接
    assert "100" in lines[1]          # 数字取首条


def test_merge_join_unmatched_fill_empty(tmp_path):
    """left join 未匹配行右表列填空"""
    left = tmp_path / "l.csv"
    left.write_text("k,v\n1,a\n2,b\n9,z\n", encoding="utf-8")
    right = tmp_path / "r.csv"
    right.write_text("k,w\n1,x\n2,y\n", encoding="utf-8")
    out = tmp_path / "out.xlsx"
    out_text = run("merge", str(left), str(right), "--on", "k", "--how", "left", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    assert len(rows) == 4  # 表头 + 3 行
    d = {r[0]: r for r in rows[1:]}
    assert d["9"][2] in (None, "")  # k=9 未匹配，w 列填空
    assert "未匹配：1" in out_text
    assert "匹配率：66.7%" in out_text


def test_merge_join_key_not_found(tmp_path):
    """key 列不存在 → 友好报错 exit 1"""
    left = tmp_path / "l.csv"
    left.write_text("k,v\n1,a\n", encoding="utf-8")
    right = tmp_path / "r.csv"
    right.write_text("k,w\n1,x\n", encoding="utf-8")
    out = tmp_path / "out.csv"
    r = subprocess.run([PY, str(SRC), "merge", str(left), str(right), "--on", "不存在", "-o", str(out)],
                       capture_output=True, text=True)
    assert r.returncode == 1
    assert "关键列" in r.stdout


def test_merge_join_then_pivot(tmp_path):
    """管道组合：merge → pivot"""
    left = tmp_path / "l.csv"
    left.write_text("ID,区域\n1,华东\n2,华北\n1,华东\n", encoding="utf-8")
    right = tmp_path / "r.csv"
    right.write_text("ID,金额\n1,100\n2,200\n", encoding="utf-8")
    mid = tmp_path / "mid.csv"
    run("merge", str(left), str(right), "--on", "ID", "-o", str(mid))
    out = tmp_path / "pv.xlsx"
    run("pivot", str(mid), "--rows", "区域", "--cols", "ID", "--values", "金额:sum", "-o", str(out))
    import openpyxl
    rows = list(openpyxl.load_workbook(out).active.values)
    assert rows[0][0] == "区域"
    assert len(rows) >= 3  # 表头 + 至少 2 个区域


def test_merge_join_then_diff(tmp_path):
    """管道组合：merge → diff"""
    left = tmp_path / "l.csv"
    left.write_text("ID,姓名\n1,张三\n2,李四\n", encoding="utf-8")
    right = tmp_path / "r.csv"
    right.write_text("ID,工资\n1,8000\n2,12000\n", encoding="utf-8")
    merged = tmp_path / "merged.csv"
    run("merge", str(left), str(right), "--on", "ID", "-o", str(merged))
    expected = tmp_path / "expected.csv"
    expected.write_text("ID,姓名,工资\n1,张三,8000\n2,李四,12000\n", encoding="utf-8")
    out = run("diff", str(merged), str(expected), "--key", "ID")
    assert "新增：0" in out
    assert "删除：0" in out
    assert "修改：0" in out

# ---------- sql 命令（v0.11） ----------
def _sql_sample(tmp_path):
    p = tmp_path / "数据.csv"
    p.write_text(
        "部门,金额,日期\n"
        "销售,100,2026-01-01\n"
        "技术,200,2026-02-01\n"
        "销售,150,2026-03-01\n"
        "技术,300,2026-01-15\n"
        "行政,80,2026-04-01\n",
        encoding="utf-8")
    return p


def test_sql_basic_select(tmp_path):
    """基本 SELECT 查询"""
    src = _sql_sample(tmp_path)
    out = run("sql", "SELECT * FROM data", "-i", str(src))
    assert "部门" in out and "金额" in out
    assert "销售" in out and "技术" in out
    assert "共 5 行" in out


def test_sql_chinese_columns(tmp_path):
    """中文列名查询"""
    src = _sql_sample(tmp_path)
    out = run("sql", "SELECT 部门, 金额 FROM data WHERE 金额 > 100", "-i", str(src))
    assert "技术" in out
    assert "行政" not in out  # 80 < 100


def test_sql_aggregation(tmp_path):
    """聚合函数 SUM/COUNT/AVG"""
    src = _sql_sample(tmp_path)
    out = run("sql", "SELECT COUNT(*) AS 数量, SUM(金额) AS 总额 FROM data", "-i", str(src))
    assert "5" in out  # 5 行
    assert "830" in out  # 100+200+150+300+80=830


def test_sql_group_by(tmp_path):
    """GROUP BY 分组聚合"""
    src = _sql_sample(tmp_path)
    out = run("sql", "SELECT 部门, SUM(金额) AS 总额 FROM data GROUP BY 部门 ORDER BY 总额 DESC", "-i", str(src))
    lines = out.strip().split("\n")
    # 技术 500 > 销售 250 > 行政 80
    assert "技术" in lines[1]
    assert "500" in lines[1]
    assert "销售" in lines[2]
    assert "250" in lines[2]


def test_sql_type_inference(tmp_path):
    """类型推断：数值列可做数值比较"""
    src = _sql_sample(tmp_path)
    out = run("sql", "SELECT 部门, 金额 FROM data WHERE 金额 >= 200", "-i", str(src))
    assert "技术" in out
    assert "销售" not in out  # 100, 150 都 < 200


def test_sql_reject_non_select(tmp_path):
    """拒绝非 SELECT 语句"""
    src = _sql_sample(tmp_path)
    r = subprocess.run([PY, str(SRC), "sql", "INSERT INTO data VALUES (1,2,3)", "-i", str(src)],
                       capture_output=True, text=True)
    assert r.returncode == 1
    assert "只支持 SELECT" in r.stdout


def test_sql_syntax_error(tmp_path):
    """SQL 语法错误友好提示"""
    src = _sql_sample(tmp_path)
    r = subprocess.run([PY, str(SRC), "sql", "SELECTT * FORM data", "-i", str(src)],
                       capture_output=True, text=True)
    assert r.returncode == 1
    assert "SQL" in r.stdout


def test_sql_column_not_found(tmp_path):
    """列名不存在时友好提示并列出可用列"""
    src = _sql_sample(tmp_path)
    r = subprocess.run([PY, str(SRC), "sql", "SELECT 不存在的列 FROM data", "-i", str(src)],
                       capture_output=True, text=True)
    assert r.returncode == 1
    assert "列名不存在" in r.stdout or "no such column" in r.stdout.lower()
    assert "部门" in r.stdout  # 列出可用列名


def test_sql_output_to_file(tmp_path):
    """输出到 xlsx 文件"""
    src = _sql_sample(tmp_path)
    out_xlsx = tmp_path / "结果.xlsx"
    run("sql", "SELECT 部门, SUM(金额) AS 总额 FROM data GROUP BY 部门", "-i", str(src), "-o", str(out_xlsx))
    import openpyxl
    rows = list(openpyxl.load_workbook(out_xlsx).active.values)
    assert rows[0] == ("部门", "总额")
    assert len(rows) == 4  # 表头 + 3 个部门


def test_sql_print_preview(tmp_path):
    """--print 预览前 N 行"""
    src = _sql_sample(tmp_path)
    out = run("sql", "SELECT * FROM data", "-i", str(src), "--print", "2")
    lines = out.strip().split("\n")
    # 表头 + 2 行数据 + 共 X 行提示
    assert "共 5 行" in out
    assert "已显示前 2 行" in out


def test_sql_where_date(tmp_path):
    """WHERE 条件过滤日期"""
    src = _sql_sample(tmp_path)
    out = run("sql", "SELECT 部门, 日期 FROM data WHERE 日期 >= '2026-02-01'", "-i", str(src))
    assert "技术" in out  # 2026-02-01
    assert "销售" in out  # 2026-03-01
    assert "行政" in out  # 2026-04-01
    lines = [l for l in out.strip().split("\n") if l and not l.startswith("部门") and not l.startswith("共")]
    assert len(lines) == 3




def test_sql_multi_sheet(tmp_path):
    """多 sheet xlsx 用 --sheet 指定"""
    import openpyxl
    xlsx = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["部门", "金额"])
    ws1.append(["销售", 100])
    ws2 = wb.create_sheet("数据")
    ws2.append(["部门", "金额"])
    ws2.append(["技术", 200])
    ws2.append(["技术", 300])
    wb.save(xlsx)
    # 默认读第一个 sheet
    out1 = run("sql", "SELECT SUM(金额) FROM data", "-i", str(xlsx))
    assert "100" in out1
    # --sheet 指定
    out2 = run("sql", "SELECT SUM(金额) FROM data", "-i", str(xlsx), "--sheet", "数据")
    assert "500" in out2

def test_sql_with_cte(tmp_path):
    """WITH (CTE) 语句支持"""
    src = _sql_sample(tmp_path)
    out = run("sql", "WITH t AS (SELECT 部门, SUM(金额) AS 总额 FROM data GROUP BY 部门) SELECT * FROM t WHERE 总额 > 100", "-i", str(src))
    assert "技术" in out
    assert "销售" in out
    assert "行政" not in out  # 80 < 100\n

# ========== chart 命令测试 ==========

def _chart_sample(tmp_path):
    """图表测试用 CSV"""
    p = tmp_path / "sales.csv"
    p.write_text(
        "部门,金额,数量\n"
        "销售,100,10\n"
        "技术,200,20\n"
        "行政,80,5\n"
        "市场,150,15\n",
        encoding="utf-8")
    return p


def _chart_pie_many(tmp_path):
    """超过 8 类的饼图测试数据"""
    p = tmp_path / "many.csv"
    lines = ["类别,值"]
    for i in range(12):
        lines.append(f"类别{i},{(12 - i) * 10}")
    p.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return p


def test_chart_bar(tmp_path):
    """bar 柱状图出图"""
    src = _chart_sample(tmp_path)
    out = tmp_path / "bar.png"
    run("chart", "-i", str(src), "--type", "bar", "--x", "部门", "--y", "金额", "--out", str(out))
    assert out.exists()
    assert out.stat().st_size > 10000  # >10KB


def test_chart_line(tmp_path):
    """line 折线图出图"""
    src = _chart_sample(tmp_path)
    out = tmp_path / "line.png"
    run("chart", "-i", str(src), "--type", "line", "--x", "部门", "--y", "金额,数量", "--out", str(out))
    assert out.exists()
    assert out.stat().st_size > 10000


def test_chart_pie(tmp_path):
    """pie 饼图出图"""
    src = _chart_sample(tmp_path)
    out = tmp_path / "pie.png"
    run("chart", "-i", str(src), "--type", "pie", "--x", "部门", "--y", "金额", "--out", str(out))
    assert out.exists()
    assert out.stat().st_size > 10000


def test_chart_scatter(tmp_path):
    """scatter 散点图出图"""
    src = _chart_sample(tmp_path)
    out = tmp_path / "scatter.png"
    run("chart", "-i", str(src), "--type", "scatter", "--x", "金额", "--y", "数量", "--out", str(out))
    assert out.exists()
    assert out.stat().st_size > 10000


def test_chart_sql(tmp_path):
    """--sql 模式出图"""
    src = _chart_sample(tmp_path)
    out = tmp_path / "sql_chart.png"
    run("chart", "-i", str(src), "--type", "bar",
        "--sql", "SELECT 部门, SUM(金额) AS 总额 FROM data GROUP BY 部门 ORDER BY 总额 DESC",
        "--out", str(out))
    assert out.exists()
    assert out.stat().st_size > 10000


def test_chart_title(tmp_path):
    """--title 中文标题"""
    src = _chart_sample(tmp_path)
    out = tmp_path / "titled.png"
    r = run("chart", "-i", str(src), "--type", "bar", "--x", "部门", "--y", "金额",
            "--title", "各部门销售额", "--out", str(out))
    assert out.exists()
    assert "图表已生成" in r


def test_chart_pie_merge(tmp_path):
    """pie 超过 8 类自动归并「其他」"""
    src = _chart_pie_many(tmp_path)
    out = tmp_path / "pie_many.png"
    r = run("chart", "-i", str(src), "--type", "pie", "--x", "类别", "--y", "值", "--out", str(out))
    assert out.exists()
    assert out.stat().st_size > 10000


def test_chart_default_out(tmp_path):
    """默认输出 chart.png"""
    src = _chart_sample(tmp_path)
    r = run("chart", "-i", str(src), "--type", "bar", "--x", "部门", "--y", "金额", cwd=str(tmp_path))
    assert (tmp_path / "chart.png").exists()
    assert "chart.png" in r


def test_chart_column_not_found(tmp_path):
    """列不存在时报错"""
    src = _chart_sample(tmp_path)
    out = tmp_path / "err.png"
    r = subprocess.run([PY, str(SRC), "chart", "-i", str(src), "--type", "bar",
                        "--x", "不存在的列", "--y", "金额", "--out", str(out)],
                       capture_output=True, text=True)
    assert r.returncode == 1
    assert "不存在" in r.stdout


def test_chart_file_not_found(tmp_path):
    """文件不存在时报错"""
    r = subprocess.run([PY, str(SRC), "chart", "-i", str(tmp_path / "nope.csv"),
                        "--type", "bar", "--x", "a", "--y", "b"],
                       capture_output=True, text=True)
    assert r.returncode == 1


def test_chart_scatter_non_numeric_x(tmp_path):
    """scatter 图 x 列非数值时报错"""
    src = _chart_sample(tmp_path)
    out = tmp_path / "err2.png"
    r = subprocess.run([PY, str(SRC), "chart", "-i", str(src), "--type", "scatter",
                        "--x", "部门", "--y", "金额", "--out", str(out)],
                       capture_output=True, text=True)
    assert r.returncode == 1
    assert "数值" in r.stdout


def test_chart_no_args(tmp_path):
    """缺少 --x/--y 和 --sql 时报错"""
    src = _chart_sample(tmp_path)
    r = subprocess.run([PY, str(SRC), "chart", "-i", str(src), "--type", "bar"],
                       capture_output=True, text=True)
    assert r.returncode == 1
    assert "--sql" in r.stdout or "--x" in r.stdout

# ========== template 命令测试 ==========

def _make_data_xlsx(path, rows=None):
    """造数据表 xlsx"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "班级", "分数"])
    if rows is None:
        rows = [["张三", "一班", 95], ["李四", "二班", 88], ["王五", "一班", 100.0]]
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()


def _make_docx_tpl(path, texts=None):
    """造 docx 模板"""
    from docx import Document
    doc = Document()
    if texts is None:
        texts = ["亲爱的{{姓名}}家长：", "您的孩子在{{班级}}获得了{{分数}}分。"]
    for t in texts:
        doc.add_paragraph(t)
    doc.save(path)


def _make_xlsx_tpl(path, cells=None):
    """造 xlsx 模板"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    if cells is None:
        cells = {"A1": "学生：{{姓名}}", "B1": "成绩：{{分数}}"}
    for coord, val in cells.items():
        ws[coord] = val
    wb.save(path)
    wb.close()


def test_template_docx_basic(tmp_path):
    """docx 模板基本替换"""
    data = tmp_path / "data.xlsx"
    tpl = tmp_path / "tpl.docx"
    out = tmp_path / "output"
    _make_data_xlsx(data, [["张三", "一班", 95]])
    _make_docx_tpl(tpl)
    r = subprocess.run([PY, str(SRC), "template", str(data), "--tpl", str(tpl), "--out", str(out), "--name-col", "姓名"],
                       capture_output=True, text=True)
    assert r.returncode == 0
    assert (out / "张三.docx").exists()
    from docx import Document
    doc = Document(str(out / "张三.docx"))
    text = " ".join(p.text for p in doc.paragraphs)
    assert "张三" in text and "一班" in text and "95" in text


def test_template_xlsx_basic(tmp_path):
    """xlsx 模板基本替换"""
    data = tmp_path / "data.xlsx"
    tpl = tmp_path / "tpl.xlsx"
    out = tmp_path / "output"
    _make_data_xlsx(data, [["张三", "一班", 95]])
    _make_xlsx_tpl(tpl)
    r = subprocess.run([PY, str(SRC), "template", str(data), "--tpl", str(tpl), "--out", str(out), "--name-col", "姓名"],
                       capture_output=True, text=True)
    assert r.returncode == 0
    assert (out / "张三.xlsx").exists()
    import openpyxl
    wb = openpyxl.load_workbook(str(out / "张三.xlsx"))
    ws = wb.active
    assert ws["A1"].value == "学生：张三"
    assert ws["B1"].value == "成绩：95"
    wb.close()


def test_template_cross_run_placeholder(tmp_path):
    """跨 run 拆分的占位符也能替换"""
    from docx import Document
    data = tmp_path / "data.xlsx"
    tpl = tmp_path / "tpl.docx"
    out = tmp_path / "output"
    _make_data_xlsx(data, [["张三", "一班", 95]])
    # 手动造跨 run 拆分
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("{{")
    para.add_run("姓名")
    para.add_run("}}")
    para.add_run(" 同学")
    doc.save(str(tpl))
    r = subprocess.run([PY, str(SRC), "template", str(data), "--tpl", str(tpl), "--out", str(out)],
                       capture_output=True, text=True)
    assert r.returncode == 0
    files = list(out.iterdir())
    assert len(files) == 1
    doc2 = Document(str(files[0]))
    assert "张三" in doc2.paragraphs[0].text


def test_template_filename_tpl(tmp_path):
    """--filename-tpl 自定义文件名"""
    data = tmp_path / "data.xlsx"
    tpl = tmp_path / "tpl.docx"
    out = tmp_path / "output"
    _make_data_xlsx(data, [["张三", "一班", 95]])
    _make_docx_tpl(tpl)
    r = subprocess.run([PY, str(SRC), "template", str(data), "--tpl", str(tpl), "--out", str(out),
                        "--filename-tpl", "{姓名}_{班级}"],
                       capture_output=True, text=True)
    assert r.returncode == 0
    assert (out / "张三_一班.docx").exists()


def test_template_filename_sanitize(tmp_path):
    """文件名非法字符自动清洗"""
    data = tmp_path / "data.xlsx"
    tpl = tmp_path / "tpl.docx"
    out = tmp_path / "output"
    _make_data_xlsx(data, [["张/三", "一班", 95]])
    _make_docx_tpl(tpl, ["{{姓名}}"])
    r = subprocess.run([PY, str(SRC), "template", str(data), "--tpl", str(tpl), "--out", str(out), "--name-col", "姓名"],
                       capture_output=True, text=True)
    assert r.returncode == 0
    files = list(out.iterdir())
    assert len(files) == 1
    assert "/" not in files[0].name


def test_template_duplicate_names(tmp_path):
    """重名自动加序号"""
    data = tmp_path / "data.xlsx"
    tpl = tmp_path / "tpl.docx"
    out = tmp_path / "output"
    _make_data_xlsx(data, [["张三", "一班", 95], ["张三", "二班", 88]])
    _make_docx_tpl(tpl, ["{{姓名}} {{班级}}"])
    r = subprocess.run([PY, str(SRC), "template", str(data), "--tpl", str(tpl), "--out", str(out), "--name-col", "姓名"],
                       capture_output=True, text=True)
    assert r.returncode == 0
    files = sorted(out.iterdir())
    assert len(files) == 2
    names = [f.name for f in files]
    assert "张三.docx" in names
    assert "张三_1.docx" in names


def test_template_number_format(tmp_path):
    """数字格式化：1.0 -> 1"""
    data = tmp_path / "data.xlsx"
    tpl = tmp_path / "tpl.docx"
    out = tmp_path / "output"
    _make_data_xlsx(data, [["张三", "一班", 100.0]])
    _make_docx_tpl(tpl, ["{{分数}}"])
    r = subprocess.run([PY, str(SRC), "template", str(data), "--tpl", str(tpl), "--out", str(out)],
                       capture_output=True, text=True)
    assert r.returncode == 0
    from docx import Document
    doc = Document(str(list(out.iterdir())[0]))
    assert "100" in doc.paragraphs[0].text
    assert "100.0" not in doc.paragraphs[0].text


def test_template_missing_col_warning(tmp_path):
    """缺失列占位符警告但成功"""
    data = tmp_path / "data.xlsx"
    tpl = tmp_path / "tpl.docx"
    out = tmp_path / "output"
    _make_data_xlsx(data, [["张三", "一班", 95]])
    _make_docx_tpl(tpl, ["{{姓名}} {{不存在的列}}"])
    r = subprocess.run([PY, str(SRC), "template", str(data), "--tpl", str(tpl), "--out", str(out)],
                       capture_output=True, text=True)
    assert r.returncode == 0
    assert "警告" in r.stdout or "⚠" in r.stdout


def test_template_missing_col_strict(tmp_path):
    """--strict 模式缺失列报错"""
    data = tmp_path / "data.xlsx"
    tpl = tmp_path / "tpl.docx"
    out = tmp_path / "output"
    _make_data_xlsx(data, [["张三", "一班", 95]])
    _make_docx_tpl(tpl, ["{{姓名}} {{不存在的列}}"])
    r = subprocess.run([PY, str(SRC), "template", str(data), "--tpl", str(tpl), "--out", str(out), "--strict"],
                       capture_output=True, text=True)
    assert r.returncode == 1


def test_template_dry_run(tmp_path):
    """--dry-run 预览不写盘"""
    data = tmp_path / "data.xlsx"
    tpl = tmp_path / "tpl.docx"
    out = tmp_path / "output"
    _make_data_xlsx(data)
    _make_docx_tpl(tpl)
    r = subprocess.run([PY, str(SRC), "template", str(data), "--tpl", str(tpl), "--out", str(out), "--dry-run"],
                       capture_output=True, text=True)
    assert r.returncode == 0
    assert "预览" in r.stdout
    assert not out.exists() or len(list(out.iterdir())) == 0


def test_template_data_not_found(tmp_path):
    """数据表不存在报错"""
    tpl = tmp_path / "tpl.docx"
    _make_docx_tpl(tpl)
    r = subprocess.run([PY, str(SRC), "template", str(tmp_path / "nope.xlsx"), "--tpl", str(tpl), "--out", str(tmp_path / "out")],
                       capture_output=True, text=True)
    assert r.returncode == 1


def test_template_tpl_not_found(tmp_path):
    """模板不存在报错"""
    data = tmp_path / "data.xlsx"
    _make_data_xlsx(data)
    r = subprocess.run([PY, str(SRC), "template", str(data), "--tpl", str(tmp_path / "nope.docx"), "--out", str(tmp_path / "out")],
                       capture_output=True, text=True)
    assert r.returncode == 1


def test_template_no_placeholders_warning(tmp_path):
    """模板无占位符时警告但 exit 0"""
    data = tmp_path / "data.xlsx"
    tpl = tmp_path / "tpl.docx"
    out = tmp_path / "output"
    _make_data_xlsx(data, [["张三", "一班", 95]])
    from docx import Document
    doc = Document()
    doc.add_paragraph("没有占位符的模板")
    doc.save(str(tpl))
    r = subprocess.run([PY, str(SRC), "template", str(data), "--tpl", str(tpl), "--out", str(out)],
                       capture_output=True, text=True)
    assert r.returncode == 0
    assert "警告" in r.stdout or "⚠" in r.stdout


def test_template_xlsx_formula_preserved(tmp_path):
    """xlsx 模板公式不被替换"""
    data = tmp_path / "data.xlsx"
    tpl = tmp_path / "tpl.xlsx"
    out = tmp_path / "output"
    _make_data_xlsx(data, [["张三", "一班", 95]])
    _make_xlsx_tpl(tpl, {"A1": "{{姓名}}", "B1": "=1+1"})
    r = subprocess.run([PY, str(SRC), "template", str(data), "--tpl", str(tpl), "--out", str(out)],
                       capture_output=True, text=True)
    assert r.returncode == 0
    import openpyxl
    wb = openpyxl.load_workbook(str(list(out.iterdir())[0]))
    ws = wb.active
    assert ws["A1"].value == "张三"
    assert ws["B1"].value == "=1+1"
    wb.close()



# ---------- fill 空值向下填充 ----------
def _make_fill_csv(path):
    path.write_text("部门,姓名,金额\n销售,张三,100\n,李四,\n,王五,200\n技术,赵六,\n,钱七,300\n", encoding="utf-8")


def test_fill_basic(tmp_path):
    src = tmp_path / "a.csv"
    _make_fill_csv(src)
    out = tmp_path / "b.csv"
    run("fill", str(src), "-o", str(out))
    lines = out.read_text(encoding="utf-8-sig").strip().splitlines()
    assert lines[0] == "部门,姓名,金额"
    assert lines[1] == "销售,张三,100"
    assert lines[2] == "销售,李四,100"
    assert lines[3] == "销售,王五,200"
    assert lines[4] == "技术,赵六,200"
    assert lines[5] == "技术,钱七,300"


def test_fill_cols_by_name(tmp_path):
    src = tmp_path / "a.csv"
    _make_fill_csv(src)
    out = tmp_path / "b.csv"
    run("fill", str(src), "-o", str(out), "--cols", "部门")
    lines = out.read_text(encoding="utf-8-sig").strip().splitlines()
    assert lines[2] == "销售,李四,"
    assert lines[3] == "销售,王五,200"


def test_fill_cols_by_letter(tmp_path):
    src = tmp_path / "a.csv"
    _make_fill_csv(src)
    out = tmp_path / "b.csv"
    run("fill", str(src), "-o", str(out), "--cols", "A")
    lines = out.read_text(encoding="utf-8-sig").strip().splitlines()
    assert lines[2] == "销售,李四,"
    assert lines[4] == "技术,赵六,"


def test_fill_unmerge_xlsx(tmp_path):
    import openpyxl
    src = tmp_path / "merged.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["部门", "姓名", "金额"])
    ws.append(["销售", "张三", 100])
    ws.append(["", "李四", 200])
    ws.append(["", "王五", 300])
    ws.merge_cells("A2:A4")
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("fill", str(src), "-o", str(out), "--unmerge")
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    assert ws2.cell(2, 1).value == "销售"
    assert ws2.cell(3, 1).value == "销售"
    assert ws2.cell(4, 1).value == "销售"
    assert ws2.cell(3, 3).value == 200
    assert len(list(ws2.merged_cells.ranges)) == 0
    wb2.close()


def test_fill_group_by_no_cross_contamination(tmp_path):
    src = tmp_path / "a.csv"
    src.write_text("部门,金额\n销售,100\n,\n技术,200\n,\n", encoding="utf-8")
    out = tmp_path / "b.csv"
    run("fill", str(src), "-o", str(out), "--group-by", "部门")
    lines = out.read_text(encoding="utf-8-sig").strip().splitlines()
    assert lines[1] == "销售,100"
    assert lines[2] == "销售,100"
    assert lines[3] == "技术,200"
    assert lines[4] == "技术,200"


def test_fill_min_gap(tmp_path):
    src = tmp_path / "a.csv"
    src.write_text("v\n1\n\n\n2\n\n\n\n3\n", encoding="utf-8")
    out = tmp_path / "b.csv"
    run("fill", str(src), "-o", str(out), "--min-gap", "2")
    lines = out.read_text(encoding="utf-8-sig").strip().splitlines()
    assert lines[1] == "1"
    assert lines[2] == "1"
    assert lines[3] == "1"
    assert lines[4] == "2"
    assert lines[5] == "\"\""
    assert lines[6] == "\"\""
    assert lines[7] == "\"\""
    assert lines[8] == "3"


def test_fill_report_no_output(tmp_path):
    src = tmp_path / "a.csv"
    _make_fill_csv(src)
    out = run("fill", str(src), "--report")
    assert "剩余空值" in out
    assert not (tmp_path / "b.csv").exists()


def test_fill_header_not_filled(tmp_path):
    src = tmp_path / "a.csv"
    src.write_text("部门,金额\n销售,100\n,\n", encoding="utf-8")
    out = tmp_path / "b.csv"
    run("fill", str(src), "-o", str(out), "--cols", "部门")
    lines = out.read_text(encoding="utf-8-sig").strip().splitlines()
    assert lines[0] == "部门,金额"
    assert lines[1] == "销售,100"
    assert lines[2] == "销售,"


def test_fill_all_empty_column(tmp_path):
    src = tmp_path / "a.csv"
    src.write_text("a,b\n,1\n,2\n,3\n", encoding="utf-8")
    out = tmp_path / "b.csv"
    run("fill", str(src), "-o", str(out))
    lines = out.read_text(encoding="utf-8-sig").strip().splitlines()
    assert lines[0] == "a,b"
    assert lines[1] == ",1"
    assert lines[2] == ",2"
    assert lines[3] == ",3"


def test_fill_preserves_types_xlsx(tmp_path):
    import openpyxl
    src = tmp_path / "types.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["数值", "日期"])
    ws.append([1.5, "2026-01-01"])
    ws.append(["", ""])
    ws.append([3.0, "2026-03-01"])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("fill", str(src), "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    assert ws2.cell(2, 1).value == 1.5
    assert ws2.cell(3, 1).value == 1.5
    assert isinstance(ws2.cell(3, 1).value, float)
    assert ws2.cell(4, 1).value == 3.0
    wb2.close()


def test_fill_multi_sheet(tmp_path):
    import openpyxl
    src = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["A", "B"])
    ws1.append(["x", ""])
    ws1.append(["", "y"])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["C", "D"])
    ws2.append(["m", ""])
    ws2.append(["", "n"])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("fill", str(src), "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws1 = wb2["Sheet1"]
    assert ws1.cell(2, 1).value == "x"
    assert ws1.cell(3, 1).value == "x"
    ws2 = wb2["Sheet2"]
    assert ws2.cell(2, 1).value == "m"
    assert ws2.cell(3, 1).value == "m"
    wb2.close()


def test_fill_csv_unmerge_hint(tmp_path):
    src = tmp_path / "a.csv"
    src.write_text("a,b\n1,2\n,3\n", encoding="utf-8")
    out = tmp_path / "b.csv"
    r = subprocess.run([PY, str(SRC), "fill", str(src), "-o", str(out), "--unmerge"],
                       capture_output=True, text=True)
    assert r.returncode == 0
    assert "CSV" in r.stdout and "合并单元格" in r.stdout


def test_fill_file_not_found(tmp_path):
    r = subprocess.run([PY, str(SRC), "fill", str(tmp_path / "nope.csv"), "-o", str(tmp_path / "b.csv")],
                       capture_output=True, text=True)
    assert r.returncode == 1
    assert "不存在" in r.stdout


def test_fill_bad_column(tmp_path):
    src = tmp_path / "a.csv"
    src.write_text("a,b\n1,2\n,3\n", encoding="utf-8")
    r = subprocess.run([PY, str(SRC), "fill", str(src), "-o", str(tmp_path / "b.csv"), "--cols", "不存在的列"],
                       capture_output=True, text=True)
    assert r.returncode == 1
    assert "找不到列" in r.stdout



# --- Regression tests for v0.14.1 bug fixes ---

def test_fill_group_by_first_group_merged(tmp_path):
    """Bug1: --group-by with merged cells must fill first group too."""
    import openpyxl
    src = tmp_path / "merged_group.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["部门", "姓名", "金额"])
    ws.append(["华东", "张三", 100])
    ws.append(["华东", "李四", None])
    ws.append(["华东", "王五", 300])
    ws.append(["华北", "赵六", 200])
    ws.append(["华北", "钱七", None])
    ws.append(["华北", "孙八", 400])
    ws.merge_cells("A2:A4")
    ws.merge_cells("A5:A7")
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("fill", str(src), "-o", str(out), "--unmerge", "--group-by", "部门")
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    assert ws2.cell(2, 1).value == "华东"
    assert ws2.cell(3, 1).value == "华东"
    assert ws2.cell(3, 3).value == 100, "first group empty cell should be filled from first group value"
    assert ws2.cell(4, 3).value == 300
    assert ws2.cell(5, 1).value == "华北"
    assert ws2.cell(6, 1).value == "华北"
    assert ws2.cell(6, 3).value == 200, "second group empty cell should be filled"
    assert ws2.cell(7, 3).value == 400
    wb2.close()


def test_fill_group_by_repeated_values_no_merge(tmp_path):
    """Bug1b: --group-by with plain repeated values (no merge) must fill within groups."""
    import openpyxl
    src = tmp_path / "repeated_group.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["部门", "姓名", "金额"])
    ws.append(["华东", "张三", 100])
    ws.append(["华东", "李四", None])
    ws.append(["华东", "王五", None])
    ws.append(["华北", "赵六", 200])
    ws.append(["华北", "钱七", None])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("fill", str(src), "-o", str(out), "--group-by", "部门")
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    assert ws2.cell(3, 3).value == 100, "first group repeated-value empty should fill"
    assert ws2.cell(4, 3).value == 100, "first group repeated-value empty should fill"
    assert ws2.cell(6, 3).value == 200, "second group repeated-value empty should fill"
    wb2.close()


def test_fill_merged_cells_without_unmerge_no_crash(tmp_path):
    """Bug2: fill without --unmerge on merged cells must not crash."""
    import openpyxl
    src = tmp_path / "merged_no_unmerge.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["部门", "金额"])
    ws.append(["销售", 100])
    ws.append([None, None])
    ws.append(["技术", 200])
    ws.merge_cells("A2:A3")
    wb.save(src)
    out = tmp_path / "out.xlsx"
    r = subprocess.run([PY, str(SRC), "fill", str(src), "-o", str(out)],
                       capture_output=True, text=True)
    assert r.returncode == 0, f"Should not crash: {r.stderr}"
    assert "合并单元格" in r.stderr or "合并单元格" in r.stdout


def test_fill_auto_unmerge_correctness(tmp_path):
    """Bug2b: auto-unmerge path produces correct filled values."""
    import openpyxl
    src = tmp_path / "auto_unmerge.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["部门", "金额"])
    ws.append(["销售", 100])
    ws.append([None, None])
    ws.append([None, 300])
    ws.merge_cells("A2:A4")
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("fill", str(src), "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    assert ws2.cell(2, 1).value == "销售"
    assert ws2.cell(3, 1).value == "销售"
    assert ws2.cell(4, 1).value == "销售"
    assert ws2.cell(3, 2).value == 100
    assert len(list(ws2.merged_cells.ranges)) == 0
    wb2.close()


# ==================== sort 命令测试 ====================

def test_sort_single_asc(tmp_path):
    import openpyxl
    src = tmp_path / "sort_asc.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "金额"])
    ws.append(["张三", 300])
    ws.append(["李四", 100])
    ws.append(["王五", 200])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("sort", str(src), "--by", "金额", "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert rows[0][1] == 100
    assert rows[1][1] == 200
    assert rows[2][1] == 300
    wb2.close()


def test_sort_single_desc(tmp_path):
    import openpyxl
    src = tmp_path / "sort_desc.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "金额"])
    ws.append(["张三", 100])
    ws.append(["李四", 300])
    ws.append(["王五", 200])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("sort", str(src), "--by", "金额:desc", "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert rows[0][1] == 300
    assert rows[1][1] == 200
    assert rows[2][1] == 100
    wb2.close()


def test_sort_multi_key_mixed(tmp_path):
    import openpyxl
    src = tmp_path / "sort_multi.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "部门", "金额"])
    ws.append(["A1", "B", 200])
    ws.append(["A2", "A", 100])
    ws.append(["A3", "B", 300])
    ws.append(["A4", "A", 50])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("sort", str(src), "--by", "部门:asc,金额:desc", "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert rows[0][1] == "A" and rows[0][2] == 100
    assert rows[1][1] == "A" and rows[1][2] == 50
    assert rows[2][1] == "B" and rows[2][2] == 300
    assert rows[3][1] == "B" and rows[3][2] == 200
    wb2.close()


def test_sort_numeric_awareness(tmp_path):
    import openpyxl
    src = tmp_path / "sort_num.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["编号", "值"])
    ws.append(["a", "10"])
    ws.append(["b", "2"])
    ws.append(["c", "1"])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("sort", str(src), "--by", "值", "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert rows[0][1] == "1"
    assert rows[1][1] == "2"
    assert rows[2][1] == "10"
    wb2.close()


def test_sort_thousands_separator(tmp_path):
    import openpyxl
    src = tmp_path / "sort_comma.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["项目", "金额"])
    ws.append(["A", "1,234.5"])
    ws.append(["B", "999"])
    ws.append(["C", "2,000"])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("sort", str(src), "--by", "金额", "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == "B"
    assert rows[1][0] == "A"
    assert rows[2][0] == "C"
    wb2.close()


def test_sort_natural_order(tmp_path):
    import openpyxl
    src = tmp_path / "sort_natural.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["文件名"])
    ws.append(["file10"])
    ws.append(["file2"])
    ws.append(["file1"])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("sort", str(src), "--by", "文件名", "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == "file1"
    assert rows[1][0] == "file2"
    assert rows[2][0] == "file10"
    wb2.close()


def test_sort_date_mixed(tmp_path):
    import openpyxl
    import datetime
    src = tmp_path / "sort_date.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["事件", "日期"])
    ws.append(["A", "2026-03-01"])
    ws.append(["B", datetime.date(2026, 1, 15)])
    ws.append(["C", "2026/02/10"])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("sort", str(src), "--by", "日期", "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == "B"
    assert rows[1][0] == "C"
    assert rows[2][0] == "A"
    wb2.close()


def test_sort_empty_last(tmp_path):
    import openpyxl
    src = tmp_path / "sort_empty.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "金额"])
    ws.append(["A", 100])
    ws.append(["B", None])
    ws.append(["C", 50])
    ws.append(["D", ""])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("sort", str(src), "--by", "金额", "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert rows[0][1] == 50
    assert rows[1][1] == 100
    assert rows[2][1] is None or rows[2][1] == ""
    assert rows[3][1] is None or rows[3][1] == ""
    wb2.close()


def test_sort_empty_first(tmp_path):
    import openpyxl
    src = tmp_path / "sort_ef.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "金额"])
    ws.append(["A", 100])
    ws.append(["B", None])
    ws.append(["C", 50])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("sort", str(src), "--by", "金额", "--empty-first", "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert rows[0][1] is None or rows[0][1] == ""
    assert rows[1][1] == 50
    assert rows[2][1] == 100
    wb2.close()


def test_sort_limit(tmp_path):
    import openpyxl
    src = tmp_path / "sort_limit.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "金额"])
    for i in range(10):
        ws.append(["人{}".format(i), (10 - i) * 100])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("sort", str(src), "--by", "金额:desc", "--limit", "3", "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 3
    assert rows[0][1] == 1000
    assert rows[1][1] == 900
    assert rows[2][1] == 800
    wb2.close()


def test_sort_stable(tmp_path):
    import openpyxl
    src = tmp_path / "sort_stable.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["部门", "姓名"])
    ws.append(["A", "张三"])
    ws.append(["B", "李四"])
    ws.append(["A", "王五"])
    ws.append(["B", "赵六"])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("sort", str(src), "--by", "部门", "--stable", "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.active
    rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert rows[0][1] == "张三"
    assert rows[1][1] == "王五"
    assert rows[2][1] == "李四"
    assert rows[3][1] == "赵六"
    wb2.close()


def test_sort_csv_roundtrip(tmp_path):
    import csv
    src = tmp_path / "sort.csv"
    with open(src, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["名称", "数量"])
        w.writerow(["banana", 3])
        w.writerow(["apple", 1])
        w.writerow(["cherry", 2])
    out = tmp_path / "out.csv"
    run("sort", str(src), "--by", "数量", "-o", str(out))
    with open(out, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    assert rows[0] == ["名称", "数量"]
    assert rows[1][0] == "apple"
    assert rows[2][0] == "cherry"
    assert rows[3][0] == "banana"


def test_sort_error_file_not_found(tmp_path):
    out = tmp_path / "out.xlsx"
    r = subprocess.run([PY, str(SRC), "sort", str(tmp_path / "nope.xlsx"), "--by", "列", "-o", str(out)], capture_output=True, text=True)
    assert r.returncode == 1
    assert "不存在" in r.stdout


def test_sort_error_column_not_found(tmp_path):
    import openpyxl
    src = tmp_path / "sort_err.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "金额"])
    ws.append(["A", 100])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    r = subprocess.run([PY, str(SRC), "sort", str(src), "--by", "不存在的列", "-o", str(out)], capture_output=True, text=True)
    assert r.returncode == 1
    assert "没有列" in r.stdout


def test_sort_error_limit_zero(tmp_path):
    import openpyxl
    src = tmp_path / "sort_limit_err.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "金额"])
    ws.append(["A", 100])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    r = subprocess.run([PY, str(SRC), "sort", str(src), "--by", "金额", "--limit", "0", "-o", str(out)], capture_output=True, text=True)
    assert r.returncode == 1
    assert "正整数" in r.stdout



# ---------- transpose 测试 ----------
def test_transpose_basic(tmp_path):
    import openpyxl
    src = tmp_path / "t.csv"
    src.write_text("姓名,1月,2月\n张三,100,200\n李四,300,400\n", encoding="utf-8-sig")
    out = tmp_path / "out.xlsx"
    run("transpose", str(src), "-o", str(out))
    wb = openpyxl.load_workbook(out)
    rows = list(wb.active.values)
    wb.close()
    assert list(rows[0]) == ["姓名", "张三", "李四"]
    assert list(rows[1]) == ["1月", "100", "300"]
    assert list(rows[2]) == ["2月", "200", "400"]


def test_transpose_roundtrip(tmp_path):
    src = tmp_path / "orig.csv"
    src.write_text("A,B,C\n1,2,3\n4,5,6\n", encoding="utf-8-sig")
    t1 = tmp_path / "t1.csv"
    t2 = tmp_path / "t2.csv"
    run("transpose", str(src), "-o", str(t1))
    run("transpose", str(t1), "-o", str(t2))
    orig = src.read_text(encoding="utf-8-sig")
    back = t2.read_text(encoding="utf-8-sig")
    assert orig == back


def test_transpose_no_header(tmp_path):
    src = tmp_path / "matrix.csv"
    src.write_text("1,2,3\n4,5,6\n", encoding="utf-8-sig")
    out = tmp_path / "out.csv"
    run("transpose", str(src), "--no-header", "-o", str(out))
    text = out.read_text(encoding="utf-8-sig")
    lines = text.strip().split("\n")
    assert lines[0] == "1,4"
    assert lines[1] == "2,5"
    assert lines[2] == "3,6"


def test_transpose_index_col(tmp_path):
    src = tmp_path / "wide.csv"
    src.write_text("指标,1月,2月,3月\n销售额,100,200,300\n利润,50,80,120\n", encoding="utf-8-sig")
    t1 = tmp_path / "t1.csv"
    run("transpose", str(src), "-o", str(t1))
    out = tmp_path / "idx.csv"
    run("transpose", str(t1), "--index-col", "指标", "-o", str(out))
    text = out.read_text(encoding="utf-8-sig")
    lines = text.strip().split("\n")
    assert lines[0] == "指标,销售额,利润"
    assert lines[1] == "1月,100,50"
    assert lines[2] == "2月,200,80"
    assert lines[3] == "3月,300,120"


def test_transpose_index_col_letter(tmp_path):
    src = tmp_path / "data.csv"
    src.write_text("A,B,C\n1,2,3\n4,5,6\n", encoding="utf-8-sig")
    t1 = tmp_path / "t1.csv"
    run("transpose", str(src), "-o", str(t1))
    out = tmp_path / "letter.csv"
    run("transpose", str(t1), "--index-col", "A", "-o", str(out))
    text = out.read_text(encoding="utf-8-sig")
    lines = text.strip().split("\n")
    assert lines[0] == "A,1,4"
    assert lines[1] == "B,2,5"
    assert lines[2] == "C,3,6"


def test_transpose_dup_headers(tmp_path):
    src = tmp_path / "dup.csv"
    src.write_text("A,1,2\nA,3,4\nB,5,6\n", encoding="utf-8-sig")
    out = tmp_path / "out.csv"
    run("transpose", str(src), "-o", str(out))
    text = out.read_text(encoding="utf-8-sig")
    first_line = text.strip().split("\n")[0]
    headers = first_line.split(",")
    assert headers[0] == "A"
    assert headers[1] == "A_2"
    assert headers[2] == "B"


def test_transpose_xlsx_format(tmp_path):
    import openpyxl
    src = tmp_path / "in.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["产品", "Q1", "Q2"])
    ws.append(["手机", 100, 200])
    ws.append(["电脑", 300, 400])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("transpose", str(src), "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    rows = list(wb2.active.values)
    wb2.close()
    assert list(rows[0]) == ["产品", "手机", "电脑"]
    assert list(rows[1]) == ["Q1", "100", "300"]
    assert list(rows[2]) == ["Q2", "200", "400"]


def test_transpose_multi_sheet(tmp_path):
    import openpyxl
    src = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["A", "B"])
    ws1.append([1, 2])
    ws2 = wb.create_sheet("数据")
    ws2.append(["X", "Y"])
    ws2.append([10, 20])
    ws2.append([30, 40])
    wb.save(src)
    out = tmp_path / "out.xlsx"
    run("transpose", str(src), "--sheet", "数据", "-o", str(out))
    wb2 = openpyxl.load_workbook(out)
    rows = list(wb2.active.values)
    wb2.close()
    assert list(rows[0]) == ["X", "10", "30"]
    assert list(rows[1]) == ["Y", "20", "40"]


def test_transpose_empty(tmp_path):
    src = tmp_path / "empty.csv"
    src.write_text("", encoding="utf-8-sig")
    out = tmp_path / "out.csv"
    run("transpose", str(src), "-o", str(out))


def test_transpose_single_row(tmp_path):
    src = tmp_path / "one.csv"
    src.write_text("A,B,C\n", encoding="utf-8-sig")
    out = tmp_path / "out.csv"
    run("transpose", str(src), "-o", str(out))
    text = out.read_text(encoding="utf-8-sig")
    lines = text.strip().split("\n")
    assert len(lines) == 3
    assert lines[0] == "A"
    assert lines[1] == "B"
    assert lines[2] == "C"


def test_transpose_error_file_not_found(tmp_path):
    out = tmp_path / "out.csv"
    r = subprocess.run([PY, str(SRC), "transpose", str(tmp_path / "nope.csv"), "-o", str(out)], capture_output=True, text=True)
    assert r.returncode == 1
    assert "不存在" in r.stdout


def test_transpose_error_bad_index_col(tmp_path):
    src = tmp_path / "data.csv"
    src.write_text("A,B\n1,2\n", encoding="utf-8-sig")
    t1 = tmp_path / "t1.csv"
    run("transpose", str(src), "-o", str(t1))
    out = tmp_path / "bad.csv"
    r = subprocess.run([PY, str(SRC), "transpose", str(t1), "--index-col", "不存在的列", "-o", str(out)], capture_output=True, text=True)
    assert r.returncode == 1
    assert "没有列" in r.stdout


def test_transpose_print_preview(tmp_path):
    src = tmp_path / "data.csv"
    src.write_text("姓名,成绩\n张三,90\n李四,85\n", encoding="utf-8-sig")
    out = run("transpose", str(src))
    assert "姓名" in out
    assert "张三" in out
    assert "共 1 行" in out
